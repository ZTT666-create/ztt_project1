from __future__ import annotations

import json
import sys
from collections import Counter
from pathlib import Path

import openpyxl


PATH = Path(r"D:\ztt-codex\codex-最新处理数据\policy_normalization_results_updated_已确认回写_摘要共同词已更新.xlsx")
STAGES = list(range(1998, 2026))


def terms(value):
    seen = set()
    out = []
    for raw in str(value or "").split(";"):
        item = raw.strip()
        if item and item not in seen:
            seen.add(item)
            out.append(item)
    return out


def rows(ws):
    return [list(r) for r in ws.iter_rows(values_only=True)]


def nonblank(data):
    return [r for r in data[1:] if r and r[0] not in (None, "")]


def records(ws):
    data = rows(ws)
    header = [str(v or "") for v in data[0]]
    idx = {h: i for i, h in enumerate(header)}
    out = []
    for row in data[1:]:
        doc_id = row[idx["doc_id"]]
        if doc_id in (None, ""):
            continue
        old_terms = terms(row[idx["规范关键词"]])
        final_terms = terms(row[idx["最终规范关键词"]])
        out.append({
            "doc_id": str(doc_id),
            "year": int(row[idx["year"]]),
            "title": str(row[idx["title"]] or ""),
            "old_terms": old_terms,
            "final_terms": final_terms,
            "old_count": row[idx["规范后数量"]],
            "final_count": row[idx["最终规范数量"]],
        })
    return out


def freq(records_):
    c = Counter()
    for r in records_:
        c.update(set(r["final_terms"]))
    return c


def audit(path: Path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    paper = records(wb["论文规范结果"])
    policy = records(wb["政策规范结果"])
    pc = freq(paper)
    qc = freq(policy)
    shared = set(pc) & set(qc)
    all_terms = set(pc) | set(qc)
    science_core = {t for t, n in pc.items() if n >= 3}
    policy_core = {t for t, n in qc.items() if n >= 3}
    science_keep = science_core | shared
    policy_keep = policy_core | shared

    report = {
        "status": "PASS",
        "source": {
            "sheets": wb.sheetnames,
            "paper_records": len(paper),
            "policy_records": len(policy),
            "science_unique_final": len(pc),
            "policy_unique_final": len(qc),
            "shared_final": len(shared),
        },
        "checks": {},
        "raw_summary": rows(wb["统计摘要"])[:30],
        "mapping_summary": {
            "rows": len(nonblank(rows(wb["术语映射"]))),
            "mismatched_norm_frequency_rows": 0,
            "examples": [],
        },
    }

    def check(name, expected, actual, details=None):
        ok = expected == actual
        report["checks"][name] = {"ok": ok, "expected": expected, "actual": actual}
        if details is not None:
            report["checks"][name]["details"] = details
        if not ok:
            report["status"] = "FAIL"

    # Source result sheets: counts must match their own keyword lists.
    result_mismatches = []
    for layer, rs in (("科学", paper), ("政策", policy)):
        for r in rs:
            if r["old_count"] != len(r["old_terms"]) or r["final_count"] != len(r["final_terms"]):
                result_mismatches.append({"layer": layer, "doc_id": r["doc_id"], "old_count": r["old_count"], "old_terms": len(r["old_terms"]), "final_count": r["final_count"], "final_terms": len(r["final_terms"])})
    check("规范结果内部数量", 0, len(result_mismatches), result_mismatches[:10])

    # Global fixed table.
    global_data = rows(wb["全局固定核心词表"])
    global_rows = nonblank(global_data)
    global_actual = {}
    for row in global_rows:
        term = str(row[0])
        global_actual[term] = row
    global_expected = {}
    for term in sorted(all_terms):
        s = pc.get(term, 0)
        p = qc.get(term, 0)
        global_expected[term] = (s, p, s + p, "是" if term in shared else "否", "是" if term in science_core else "否", "是" if term in policy_core else "否", "是" if term in science_keep else "否", "是" if term in policy_keep else "否")
    global_mismatches = []
    for term, expected in global_expected.items():
        row = global_actual.get(term)
        actual = None if row is None else tuple(row[1:9])
        if actual != expected:
            global_mismatches.append({"term": term, "expected": expected, "actual": actual})
    check("全局固定核心词表", len(global_expected), len(global_actual), global_mismatches[:10])
    check("全局固定核心词表字段", 0, len(global_mismatches), global_mismatches[:10])

    # Per-document fixed selection sheets.
    for sheet_name, rs, keep in (("论文固定词筛选结果", paper, science_keep), ("政策固定词筛选结果", policy, policy_keep)):
        data = rows(wb[sheet_name])
        actual_rows = nonblank(data)
        actual_by_id = {str(r[0]): r for r in actual_rows}
        mismatches = []
        for src in rs:
            selected = [t for t in src["final_terms"] if t in keep]
            expected = [src["doc_id"], src["year"], src["title"], ";".join(src["final_terms"]), ";".join(selected), len(selected), len(src["final_terms"]), len(selected) / len(src["final_terms"]) if src["final_terms"] else 0]
            actual = actual_by_id.get(src["doc_id"])
            if actual is None or actual[:8] != expected:
                mismatches.append({"doc_id": src["doc_id"], "expected": expected, "actual": actual[:8] if actual else None})
        check(sheet_name, len(rs), len(actual_rows), mismatches[:10])
        check(sheet_name + "字段内容", 0, len(mismatches), mismatches[:10])

    # Annual fixed-term summary.
    annual = rows(wb["年度固定词汇总"])
    annual_data = {int(r[0]): r for r in nonblank(annual) if isinstance(r[0], (int, float)) and int(r[0]) in STAGES}
    annual_mismatches = []
    for year in STAGES:
        p_year = [r for r in paper if r["year"] == year]
        q_year = [r for r in policy if r["year"] == year]
        p_terms = sorted({t for r in p_year for t in r["final_terms"] if t in science_keep})
        q_terms = sorted({t for r in q_year for t in r["final_terms"] if t in policy_keep})
        common = sorted(set(p_terms) & set(q_terms))
        expected = [year, len(p_year), len(p_terms), set(p_terms), len(q_year), len(q_terms), set(q_terms), len(common), set(common)]
        row = annual_data.get(year)
        actual = None if row is None else [int(row[0]), row[1], row[2], set(terms(row[3])), row[4], row[5], set(terms(row[6])), row[7], set(terms(row[8]))]
        if row is None or expected != actual:
            annual_mismatches.append({"year": year, "expected": [expected[0], expected[1], expected[2], sorted(expected[3]), expected[4], expected[5], sorted(expected[6]), expected[7], sorted(expected[8])], "actual": actual})
    check("年度固定词汇总", 28, len(annual_data), annual_mismatches[:10])
    check("年度固定词汇总字段", 0, len(annual_mismatches), annual_mismatches[:10])

    # Shared-term table: compare all nonblank rows against the final shared vocabulary.
    common_data = rows(wb["规范后共同词"])
    common_rows = nonblank(common_data)
    common_actual = {str(r[0]): tuple(r[1:4]) for r in common_rows}
    common_expected = {t: (pc[t], qc[t], pc[t] + qc[t]) for t in shared}
    common_mismatches = []
    for term, expected in common_expected.items():
        if common_actual.get(term) != expected:
            common_mismatches.append({"term": term, "expected": expected, "actual": common_actual.get(term)})
    check("规范后共同词覆盖", len(common_expected), len(common_actual), common_mismatches[:10])
    check("规范后共同词字段", 0, len(common_mismatches), common_mismatches[:10])

    # Mapping counts: compare its science/policy frequency columns when the canonical term is in final data.
    mapping = nonblank(rows(wb["术语映射"]))
    mapping_mismatches = []
    for row in mapping:
        canonical = str(row[2] or "").strip()
        if canonical in all_terms:
            expected = (pc.get(canonical, 0), qc.get(canonical, 0), pc.get(canonical, 0) + qc.get(canonical, 0))
            actual = (row[3], row[4], row[5])
            if actual != expected:
                mapping_mismatches.append({"canonical": canonical, "expected": expected, "actual": actual, "raw": row[0]})
    report["mapping_summary"] = {"rows": len(mapping), "mismatched_norm_frequency_rows": len(mapping_mismatches), "examples": mapping_mismatches[:10]}
    if mapping_mismatches:
        report["status"] = "FAIL"

    # Fixed-table rule block.
    rule_block = rows(wb["年度固定词汇总"])[1:9]
    report["rule_block"] = rule_block
    expected_rules = {
        "科学固定词表": "全时段科学频次≥3关键词 + 全时段科学-政策共享词",
        "政策固定词表": "全时段政策频次≥3关键词 + 全时段科学-政策共享词",
        "科学频次≥3关键词数": 207,
        "政策频次≥3关键词数": 158,
        "全时段科学-政策共享词数": 203,
        "科学最终固定词数": 326,
        "政策最终固定词数": 276,
    }
    rule_actual = {str(r[10]): r[11] for r in rule_block if len(r) > 11 and r[10] not in (None, "")}
    check("年度规则说明", expected_rules, rule_actual)

    sys.stdout.buffer.write(json.dumps(report, ensure_ascii=False, indent=2, default=list).encode("utf-8"))


if __name__ == "__main__":
    audit(PATH)
