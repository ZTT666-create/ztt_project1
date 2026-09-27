from pathlib import Path
import openpyxl


path = Path(r"D:\ztt-codex\codex-最新处理数据\policy_normalization_results_updated_已确认回写.xlsx")
wb = openpyxl.load_workbook(path, read_only=True, data_only=False)

approved = {}
for sheet_name in ("论文规范结果", "政策规范结果"):
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    headers = list(rows[0])
    idx = {str(v): i for i, v in enumerate(headers)}
    changed = 0
    overwritten = 0
    confirmed = 0
    count_mismatch = 0
    for row in rows[1:]:
        old = str(row[idx["规范关键词"]] or "")
        original = str(row[idx["原规范关键词"]] or "")
        final = str(row[idx["最终规范关键词"]] or "")
        if old != original:
            overwritten += 1
        if old != final:
            changed += 1
        if row[idx["审核状态"]] == "已确认":
            confirmed += 1
            for part in str(row[idx["审核备注"]] or "").split("；"):
                if "→" in part and "（" in part:
                    left, right = part.split("→", 1)
                    target = right.split("（", 1)[0]
                    approved[left] = target
        n = len([x for x in final.split(";") if x])
        if n != int(row[idx["最终规范数量"]] or 0):
            count_mismatch += 1
    print({"sheet": sheet_name, "rows": len(rows) - 1, "changed": changed, "overwritten": overwritten, "confirmed": confirmed, "count_mismatch": count_mismatch})

ws = wb["术语映射"]
rows = list(ws.iter_rows(values_only=True))
headers = list(rows[0])
idx = {str(v): i for i, v in enumerate(headers)}
dict_map = {}
for row in rows[1:]:
    raw = str(row[idx["原始词"]] or "")
    target = str(row[idx["规范词"]] or "")
    if raw:
        dict_map.setdefault(raw, set()).add(target)
missing = [k for k in approved if k not in dict_map]
mismatch = [k for k, v in approved.items() if k in dict_map and v not in dict_map[k]]
print({"dictionary_rows": len(rows) - 1, "approved_mappings": len(approved), "missing_in_dictionary": len(missing), "dictionary_target_mismatch": len(mismatch)})
print({"missing_examples": missing[:10], "mismatch_examples": mismatch[:10]})
print({"sheets": wb.sheetnames, "output_size": path.stat().st_size})
