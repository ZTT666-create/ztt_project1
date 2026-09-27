from __future__ import annotations

import csv
import itertools
import math
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook


BASE = Path(__file__).resolve().parent
DATA_DIR = BASE / "最终规范关键词_Leiden分辨率比较_1998-2025"
BOOK = BASE / "policy_normalization_results_updated_已确认回写_摘要共同词已更新.xlsx"
ASSIGNMENT_FILE = DATA_DIR / "Leiden_resolution0.3_节点社区分配_最终关键词.csv"
OUT_DIR = BASE / "最终规范关键词_桑基图社区合并_resolution0.3"

STAGES = [("A", "1998-2005"), ("B", "2006-2011"), ("C", "2012-2015"), ("D", "2016-2019"), ("E", "2020-2025")]
STAGE_BY_YEAR = {}
for stage, years in STAGES:
    a, b = years.split("-")
    for y in range(int(a), int(b) + 1):
        STAGE_BY_YEAR[y] = stage

# These terms remain in the network but are not allowed to dominate community names.
MANUAL_BACKGROUND = {
    "科技金融", "金融机构", "金融服务", "科技企业", "科技型企业", "创业投资",
    "科技创新", "金融创新", "融资", "资本市场", "金融体系", "金融发展",
}
BACKGROUND_STEMS = (
    "科技金融", "金融机构", "金融服务", "科技企业", "科技型企业", "创业投资",
    "科技创新", "金融创新", "资本市场", "金融体系", "金融发展",
)

MANUAL_NAME_OVERRIDES = {
    ("科学", "C", "3"): "中小企业质押融资",
    ("科学", "D", "1"): "空间溢出与成果转化",
    ("科学", "E", "1"): "新质生产力与效率",
    ("科学", "E", "2"): "科技信贷与财政投入",
    ("科学", "E", "3"): "金融科技监管",
    ("政策", "A", "1"): "贴息与科技中小企业",
    ("政策", "B", "1"): "科技保险与风险补偿",
    ("政策", "C", "1"): "科技保险与质押融资",
    ("政策", "C", "3"): "质押与供应链融资",
    ("政策", "D", "1"): "创投质押与风险补偿",
    ("政策", "E", "1"): "知识产权融资与科技保险",
    ("政策", "E", "3"): "数字化评估与产融对接",
}


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def write_csv(path: Path, rows: list[dict[str, object]], fields: list[str]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def keywords(value) -> list[str]:
    if value is None:
        return []
    seen = set()
    out = []
    for x in str(value).split(";"):
        x = x.strip()
        if x and x not in seen:
            seen.add(x)
            out.append(x)
    return out


def node_id(layer: str, stage: str, community: str) -> str:
    return f"{layer}_{stage}_M{community}"


def is_background(kw: str, dynamic_background: set[str]) -> bool:
    return kw in dynamic_background or kw in MANUAL_BACKGROUND or any(kw.startswith(stem) for stem in BACKGROUND_STEMS)


def main() -> None:
    assignment_rows = read_csv(ASSIGNMENT_FILE)
    assignment: dict[tuple[str, str], dict[str, dict[str, str]]] = defaultdict(dict)
    for row in assignment_rows:
        assignment[(row["Layer"], row["Stage"])][row["Label"]] = row

    # Rebuild the retained-node co-occurrence edges from the final workbook.
    docs: dict[tuple[str, str], list[list[str]]] = defaultdict(list)
    wb = load_workbook(BOOK, read_only=True, data_only=True)
    for layer, sheet in [("科学", "论文规范结果"), ("政策", "政策规范结果")]:
        ws = wb[sheet]
        headers = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
        h = {str(v): i for i, v in enumerate(headers) if v is not None}
        year_i = h["year"]
        kw_i = h["最终规范关键词"]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[year_i] is None:
                continue
            try:
                year = int(float(row[year_i]))
            except (TypeError, ValueError):
                continue
            stage = STAGE_BY_YEAR.get(year)
            kws = keywords(row[kw_i])
            if stage and kws:
                docs[(layer, stage)].append(kws)
    wb.close()

    edges: dict[tuple[str, str], Counter[tuple[str, str]]] = defaultdict(Counter)
    for key, doc_list in docs.items():
        retained = set(assignment[key])
        for doc_kws in doc_list:
            kept = sorted(retained.intersection(doc_kws))
            for a, b in itertools.combinations(kept, 2):
                edges[key][(a, b)] += 1

    # Merge every original community with fewer than 10 retained keyword nodes.
    merge_rows = []
    merged_of: dict[tuple[str, str], dict[str, str]] = defaultdict(dict)
    group_members: dict[tuple[str, str], dict[str, list[str]]] = defaultdict(lambda: defaultdict(list))
    for key, rows_by_kw in assignment.items():
        layer, stage = key
        sizes = Counter(row["Community"] for row in rows_by_kw.values())
        large = sorted([c for c, n in sizes.items() if n >= 10], key=lambda c: (-sizes[c], int(c)))
        if not large:
            large = [max(sizes, key=lambda c: (sizes[c], -int(c)))]
        candidates_by_small: dict[str, dict[str, int]] = {}
        for small, small_size in sorted(sizes.items(), key=lambda x: int(x[0])):
            if small_size >= 10:
                continue
            weights = {c: 0 for c in large}
            small_nodes = {kw for kw, row in rows_by_kw.items() if row["Community"] == small}
            for (a, b), w in edges[key].items():
                ca = rows_by_kw[a]["Community"] if a in rows_by_kw else None
                cb = rows_by_kw[b]["Community"] if b in rows_by_kw else None
                if ca == small and cb in weights:
                    weights[cb] += w
                elif cb == small and ca in weights:
                    weights[ca] += w
            candidates_by_small[small] = weights
            best = max(large, key=lambda c: (weights[c], sizes[c], -int(c)))
            fallback = "否" if weights[best] > 0 else "是（无社区间边，按最大社区兜底）"
            merge_rows.append(
                {
                    "Layer": layer,
                    "Stage": stage,
                    "OriginalCommunity": small,
                    "OriginalNodeCount": small_size,
                    "TargetOriginalCommunity": best,
                    "ConnectionWeight": weights[best],
                    "CandidateWeights": ";".join(f"{c}:{weights[c]}" for c in large),
                    "Fallback": fallback,
                }
            )
        target_of = {c: c for c in sizes}
        for small, weights in candidates_by_small.items():
            target_of[small] = max(large, key=lambda c: (weights[c], sizes[c], -int(c)))

        # Re-label merged groups by descending node count for a clean Sankey stage.
        members: dict[str, list[str]] = defaultdict(list)
        for original, target in target_of.items():
            members[target].append(original)
        ordered_targets = sorted(members, key=lambda c: (-sum(sizes[x] for x in members[c]), int(c)))
        new_id_of_target = {target: str(i + 1) for i, target in enumerate(ordered_targets)}
        for kw, row in rows_by_kw.items():
            target = target_of[row["Community"]]
            merged = new_id_of_target[target]
            merged_of[key][kw] = merged
            group_members[key][merged].append(kw)

    # Determine community-specific labels using frequency + community-level distinctiveness.
    name_rows = []
    name_lookup: dict[tuple[str, str, str], dict[str, str]] = {}
    node_rows = []
    for key, rows_by_kw in assignment.items():
        layer, stage = key
        groups = group_members[key]
        community_count = len(groups)
        community_df = Counter()
        for merged, kws in groups.items():
            community_df.update(set(kws))
        background = set(MANUAL_BACKGROUND)
        dynamic_cutoff = max(2, math.ceil(community_count * 0.5))
        background.update(k for k, n in community_df.items() if n >= dynamic_cutoff)
        for merged, kws in sorted(groups.items(), key=lambda x: int(x[0])):
            scored = []
            for kw in kws:
                row = rows_by_kw[kw]
                freq = float(row.get("Frequency") or 0)
                degree = float(row.get("Degree") or 0)
                df = community_df[kw]
                idf = math.log((1 + community_count) / (1 + df)) + 1
                score = math.log1p(freq) * idf * math.log1p(degree)
                if is_background(kw, background):
                    score *= 0.18
                scored.append((score, freq, degree, kw))
            scored.sort(key=lambda x: (-x[0], -x[1], -x[2], x[3]))
            informative = [x for x in scored if x[3] not in background]
            chosen = informative or scored
            first = chosen[0][3] if chosen else "未命名社区"
            second = chosen[1][3] if len(chosen) > 1 else ""
            if second and len(first) <= 6 and len(first) + len(second) <= 12 and first not in second and second not in first:
                name = f"{first}与{second}"
            else:
                name = first
            name_method = "区分度评分候选"
            override = MANUAL_NAME_OVERRIDES.get((layer, stage, merged))
            if override:
                name = override
                name_method = "基于区分关键词的人工校正"
            original_communities = sorted({assignment[key][kw]["Community"] for kw in kws}, key=int)
            original_sizes = Counter(assignment[key][kw]["Community"] for kw in kws)
            original_desc = ";".join(f"C{x}({original_sizes[x]})" for x in original_communities)
            top_keywords = ";".join(x[3] for x in chosen[:12])
            background_keywords = ";".join(x[3] for x in scored if is_background(x[3], background))[:2000]
            name_basis = ";".join(x[3] for x in chosen[:3])
            name_lookup[(layer, stage, merged)] = {"name": name, "original": original_desc, "top": top_keywords, "background": background_keywords, "basis": name_basis}
            name_rows.append(
                {
                    "Layer": layer, "Stage": stage, "Years": dict(STAGES)[stage], "Community": merged,
                    "NodeId": node_id(layer, stage, merged), "CommunityName": name, "NodeLabel": f"{stage}C{merged} {name}", "OriginalCommunities": original_desc,
                    "NodeCount": len(kws), "NameBasis": name_basis, "BackgroundKeywords": background_keywords,
                    "TopKeywordsByScore": top_keywords, "NameMethod": name_method,
                }
            )
            for kw in kws:
                old = rows_by_kw[kw]
                node_rows.append(
                    {
                        "Layer": layer, "Stage": stage, "Years": dict(STAGES)[stage], "Label": kw,
                        "Frequency": old["Frequency"], "Degree": old["Degree"], "OriginalCommunity": old["Community"],
                        "Community": merged, "CommunityName": name, "Resolution": "0.3", "MergeRule": "NodeCount<10→最大社区间连接权重",
                    }
                )

    flow_rows = []
    detail_rows = []
    for layer in ["科学", "政策"]:
        for (left_stage, left_years), (right_stage, right_years) in zip(STAGES, STAGES[1:]):
            left = merged_of[(layer, left_stage)]
            right = merged_of[(layer, right_stage)]
            groups: dict[tuple[str, str], list[str]] = defaultdict(list)
            for kw in sorted(set(left) & set(right)):
                groups[(left[kw], right[kw])].append(kw)
            for (src, dst), kws in sorted(groups.items(), key=lambda x: (int(x[0][0]), int(x[0][1]))):
                src_meta = name_lookup[(layer, left_stage, src)]
                dst_meta = name_lookup[(layer, right_stage, dst)]
                source = node_id(layer, left_stage, src)
                target = node_id(layer, right_stage, dst)
                flow_rows.append(
                    {
                        "Layer": layer, "FromStage": left_stage, "FromYears": left_years, "FromCommunity": src,
                        "Source": source, "SourceName": src_meta["name"], "ToStage": right_stage, "ToYears": right_years,
                        "ToCommunity": dst, "Target": target, "TargetName": dst_meta["name"], "Value": len(kws),
                        "SharedKeywordCount": len(kws), "SharedKeywords": ";".join(kws),
                    }
                )
                for kw in kws:
                    detail_rows.append(
                        {
                            "Layer": layer, "FromStage": left_stage, "FromYears": left_years, "FromCommunity": src,
                            "FromCommunityName": src_meta["name"], "ToStage": right_stage, "ToYears": right_years,
                            "ToCommunity": dst, "ToCommunityName": dst_meta["name"], "Keyword": kw,
                        }
                    )

    stats_rows = []
    for layer in ["科学", "政策"]:
        for stage, years in STAGES:
            original_sizes = Counter(row["Community"] for row in assignment[(layer, stage)].values())
            merged_sizes = Counter(merged_of[(layer, stage)].values())
            stats_rows.append(
                {
                    "Layer": layer, "Stage": stage, "Years": years,
                    "OriginalCommunityCount": len(original_sizes), "MergedCommunityCount": len(merged_sizes),
                    "OriginalSmallCommunityCount_lt10": sum(n < 10 for n in original_sizes.values()),
                    "MergedSmallCommunityCount_lt10": sum(n < 10 for n in merged_sizes.values()),
                    "OriginalMaxCommunitySize": max(original_sizes.values()), "MergedMaxCommunitySize": max(merged_sizes.values()),
                    "MergedCommunitySizes": ";".join(str(merged_sizes[x]) for x in sorted(merged_sizes, key=lambda z: int(z))),
                }
            )

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    write_csv(OUT_DIR / "合并规则审计表.csv", merge_rows, ["Layer", "Stage", "OriginalCommunity", "OriginalNodeCount", "TargetOriginalCommunity", "ConnectionWeight", "CandidateWeights", "Fallback"])
    write_csv(OUT_DIR / "合并后节点社区分配.csv", node_rows, ["Layer", "Stage", "Years", "Label", "Frequency", "Degree", "OriginalCommunity", "Community", "CommunityName", "Resolution", "MergeRule"])
    write_csv(OUT_DIR / "合并后社区命名对照表.csv", name_rows, ["Layer", "Stage", "Years", "Community", "NodeId", "CommunityName", "NodeLabel", "OriginalCommunities", "NodeCount", "NameBasis", "BackgroundKeywords", "TopKeywordsByScore", "NameMethod"])
    write_csv(OUT_DIR / "合并后桑基图社区流动.csv", flow_rows, ["Layer", "FromStage", "FromYears", "FromCommunity", "Source", "SourceName", "ToStage", "ToYears", "ToCommunity", "Target", "TargetName", "Value", "SharedKeywordCount", "SharedKeywords"])
    write_csv(OUT_DIR / "合并后桑基图关键词流动明细.csv", detail_rows, ["Layer", "FromStage", "FromYears", "FromCommunity", "FromCommunityName", "ToStage", "ToYears", "ToCommunity", "ToCommunityName", "Keyword"])
    write_csv(OUT_DIR / "合并前后社区规模统计.csv", stats_rows, ["Layer", "Stage", "Years", "OriginalCommunityCount", "MergedCommunityCount", "OriginalSmallCommunityCount_lt10", "MergedSmallCommunityCount_lt10", "OriginalMaxCommunitySize", "MergedMaxCommunitySize", "MergedCommunitySizes"])

    for r in stats_rows:
        print(f"{r['Layer']}|{r['Stage']}|{r['OriginalCommunityCount']}->{r['MergedCommunityCount']}|small {r['OriginalSmallCommunityCount_lt10']}->{r['MergedSmallCommunityCount_lt10']}")
    print(f"OUT|{OUT_DIR}")


if __name__ == "__main__":
    main()
